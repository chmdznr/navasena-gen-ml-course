"""Cek struktural deck exam prep: jumlah slide, notes, kata terlarang, panjang baris."""
import re
import sys
from pathlib import Path

from pptx import Presentation

DECK = Path(__file__).with_name("exam_prep.pptx")
N_SLIDES = 14
FACT_SLIDES = {3, 8, 9, 10, 11}
BANNED = re.compile(
    r"\bbatch\b|minggu ke|hari ke|\d+\s*(minggu|hari)\b|cakrawala|\bkalian\b|\bAnda\b"
    r"|passing score\s*[:=]?\s*\d|\d+\s*%\s*(untuk )?lulus",
    re.IGNORECASE,
)
MAX_CHARS = 90


def slide_text(slide):
    return [p.text for sh in slide.shapes if sh.has_text_frame
            for p in sh.text_frame.paragraphs if p.text.strip()]


def main():
    errs = []
    prs = Presentation(DECK)
    if len(prs.slides) != N_SLIDES:
        errs.append(f"jumlah slide {len(prs.slides)} != {N_SLIDES}")
    for i, s in enumerate(prs.slides, 1):
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        if not notes.strip():
            errs.append(f"slide {i}: notes kosong")
        if i in FACT_SLIDES and "http" not in notes:
            errs.append(f"slide {i}: notes tanpa URL sumber")
        paras = slide_text(s)
        for t in paras + [notes]:
            if m := BANNED.search(t):
                errs.append(f"slide {i}: kata terlarang {m.group(0)!r}")
        long = [p for p in paras if len(p) > MAX_CHARS and not p.startswith("http")]
        if long:
            errs.append(f"slide {i}: {len(long)} baris > {MAX_CHARS} karakter")
    print("\n".join(errs) or f"OK: {N_SLIDES} slide, notes lengkap")
    sys.exit(1 if errs else 0)


if __name__ == "__main__":
    main()
